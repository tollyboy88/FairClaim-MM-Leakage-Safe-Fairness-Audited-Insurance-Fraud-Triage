"""Create the styled Word manuscript from the LaTeX source and generated assets."""

from __future__ import annotations

from pathlib import Path
import re

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PUB = ROOT / "publication"
TEX = PUB / "manuscript" / "main.tex"
BIB = PUB / "manuscript" / "references.bib"
OUT = PUB / "manuscript" / "FairClaim-MM_manuscript.docx"
FIG = PUB / "figures"
TABLE = PUB / "tables"

# narrative_proposal preset with editorial_cover header template
NAVY = RGBColor(32, 55, 72)
BLUE = RGBColor(43, 81, 99)
GOLD = RGBColor(184, 134, 11)
GRAY = RGBColor(80, 80, 80)
MUTED = RGBColor(105, 113, 122)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_section(section, first_page=False):
    section.top_margin = Inches(0.8); section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0); section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.3); section.footer_distance = Inches(0.3)
    if first_page:
        section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.clear(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("FairClaim-MM"); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    r = p.add_run("    |    Journal manuscript"); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    for k,v in (("val","single"),("sz","4"),("space","3"),("color","D7DBE2")): bottom.set(qn(f"w:{k}"),v)
    pbdr.append(bottom); pPr.append(pbdr)
    footer = section.footer
    fp = footer.paragraphs[0]; fp.clear(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Page "); fr.font.size = Pt(8.5); fr.font.color.rgb = MUTED; add_page_field(fp)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor(35,35,35)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
    for name, size, color, before, after in (("Title",30,NAVY,0,8),("Subtitle",15,BLUE,0,10),("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,12,6),("Heading 3",12,NAVY,8,4)):
        style = doc.styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size); style.font.color.rgb = color
        style.font.bold = name != "Subtitle"; style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Equation LaTeX" not in doc.styles:
        st = doc.styles.add_style("Equation LaTeX", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Consolas"; st.font.size = Pt(9.5); st.font.color.rgb = NAVY
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; st.paragraph_format.space_before = Pt(6); st.paragraph_format.space_after = Pt(9)
    if "Figure Caption Custom" not in doc.styles:
        st = doc.styles.add_style("Figure Caption Custom", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Calibri"; st.font.size = Pt(9); st.font.italic = True; st.font.color.rgb = GRAY
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; st.paragraph_format.space_after = Pt(10); st.paragraph_format.keep_with_next = False


def bib_entries():
    text = BIB.read_text(encoding="utf-8")
    starts = list(re.finditer(r"@(\w+)\{([^,]+),", text))
    result = {}
    for i,m in enumerate(starts):
        body = text[m.end(): starts[i+1].start() if i+1 < len(starts) else len(text)]
        fields = {fm.group(1).lower(): " ".join(fm.group(2).split()).strip("{}") for fm in re.finditer(r"(\w+)\s*=\s*\{(.*?)\}(?=\s*,\s*\w+\s*=|\s*\}\s*$)", body, re.DOTALL)}
        result[m.group(2)] = fields
    return result


BIBS = bib_entries()


def cite(keys):
    chunks = []
    for key in keys.split(","):
        f = BIBS.get(key.strip(), {})
        authors = f.get("author", key).split(" and ")
        first = authors[0].strip("{} ")
        surname = first.split(",")[0] if "," in first else first.split()[-1]
        who = surname if len(authors) == 1 else (f"{surname} and {authors[1].split(',')[0]}" if len(authors) == 2 else f"{surname} et al.")
        chunks.append(f"{who}, {f.get('year','n.d.')}")
    return "(" + "; ".join(chunks) + ")"


def clean(text):
    text = re.sub(r"\\citep\{([^}]+)\}", lambda m: cite(m.group(1)), text)
    text = re.sub(r"\\textbf\{([^{}]*)\}", r"\1", text)
    text = text.replace("\\%", "%").replace("\\&", "&").replace("\\_", "_")
    text = text.replace("--", "–").replace("\\mbox{-}", "-")
    text = text.replace("\\\\", "")
    text = text.replace("$", "")
    return text.strip()


def add_table(doc, csv_name, title, max_rows=None):
    df = pd.read_csv(TABLE / csv_name)
    if max_rows: df = df.head(max_rows)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    r = p.add_run(title); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
    table = doc.add_table(rows=1, cols=len(df.columns)); table.autofit = False
    usable = Inches(7.0); width = usable / len(df.columns)
    for j,col in enumerate(df.columns):
        cell = table.rows[0].cells[j]; cell.width = width; cell.text = str(col).replace("_", " ")
        set_cell_shading(cell, "1F4E79"); set_cell_margins(cell)
        for run in cell.paragraphs[0].runs: run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(7.5)
    trPr = table.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); trPr.append(repeat)
    for i,row in df.iterrows():
        cells = table.add_row().cells
        for j,val in enumerate(row):
            if isinstance(val, float): val = f"{val:.4f}"
            cells[j].text = str(val); cells[j].width = width; cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[j]);
            if i % 2: set_cell_shading(cells[j], "F4F6F9")
            for para in cells[j].paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs: run.font.size = Pt(7.3)
        cant = OxmlElement("w:cantSplit"); table.rows[-1]._tr.get_or_add_trPr().append(cant)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_cover(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(90); p.paragraph_format.space_after = Pt(18)
    r = p.add_run("JOURNAL MANUSCRIPT"); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = GOLD
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("FairClaim-MM")
    p = doc.add_paragraph(style="Subtitle"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Leakage-Safe, Fairness-Audited Insurance Fraud Triage with Calibrated Abstention and Component-Level Multimodal Evidence")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(55)
    r = p.add_run("Author name(s), affiliation(s) and ORCID(s) to be inserted"); r.italic = True; r.font.size = Pt(10); r.font.color.rgb = GRAY
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Prepared 20 July 2026  •  Research prototype manuscript"); r.font.size = Pt(9); r.font.color.rgb = MUTED


def main():
    doc = Document(); configure_styles(doc); doc.settings.odd_and_even_pages_header_footer = False
    configure_section(doc.sections[0], first_page=True)
    add_cover(doc); doc.add_page_break()
    lines = TEX.read_text(encoding="utf-8").splitlines()
    paragraph = []
    in_doc = in_abstract = in_equation = False
    eq_lines = []
    figure_number = 0
    skip = {"\\maketitle", "\\onehalfspacing", "\\begin{document}"}

    def flush():
        nonlocal paragraph
        if paragraph:
            text = clean(" ".join(x.strip() for x in paragraph))
            if text: doc.add_paragraph(text)
            paragraph = []

    for raw in lines:
        s = raw.strip()
        if s == "\\begin{document}": in_doc = True; continue
        if not in_doc or s in skip: continue
        if s == "\\end{document}": break
        if s == "\\begin{abstract}": flush(); in_abstract = True; doc.add_heading("Abstract", level=1); continue
        if s == "\\end{abstract}": flush(); in_abstract = False; continue
        if s.startswith("\\begin{equation}") or s.startswith("\\["):
            flush(); in_equation=True; eq_lines=[]; continue
        if s.startswith("\\end{equation}") or s.startswith("\\]"):
            doc.add_paragraph(" ".join(eq_lines), style="Equation LaTeX"); in_equation=False; continue
        if in_equation: eq_lines.append(s); continue
        if s.startswith("\\section*{"):
            flush(); title = re.search(r"\{(.+)\}",s).group(1); doc.add_heading(clean(title), level=1); continue
        if s.startswith("\\section{"):
            flush(); title = re.search(r"\{(.+)\}",s).group(1); doc.add_heading(clean(title), level=1); continue
        if s.startswith("\\subsection{"):
            flush(); title = clean(re.search(r"\{(.+)\}",s).group(1)); doc.add_heading(title, level=2)
            if title.startswith("RQ1"):
                add_table(doc,"table_03_model_comparison.csv","Table 1. Model-family comparison",6)
                add_table(doc,"table_04_imbalance_comparison.csv","Table 2. Imbalance-strategy comparison",5)
            elif title.startswith("RQ2"):
                add_table(doc,"table_05_decision_policy.csv","Table 3. Held-out decision-region summary",3)
            elif title.startswith("RQ3"):
                add_table(doc,"table_06_fairness_gaps.csv","Table 4. Held-out fairness gaps",4)
                add_table(doc,"table_08_feature_importance.csv","Table 5. Leading permutation importance",6)
            elif title.startswith("Exploratory"):
                add_table(doc,"table_09_image_model.csv","Table 6. Image damage-class performance",6)
            continue
        if "\\includegraphics" in s:
            flush(); doc.add_page_break(); m = re.search(r"\{\.\./figures/([^}]+)\}",s)
            if m:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(FIG / m.group(1)), width=Inches(6.65)); figure_number += 1
            continue
        if s.startswith("\\begin{enumerate}") or s.startswith("\\end{enumerate}") or s.startswith("\\begin{figure}") or s.startswith("\\end{figure}") or s.startswith("\\centering") or s.startswith("\\bibliograph"):
            flush(); continue
        if s.startswith("\\item"):
            flush(); doc.add_paragraph(clean(s[5:].strip()), style="List Number"); continue
        if s.startswith("\\caption{"):
            flush(); cap = clean(s[len("\\caption{"):-1]); doc.add_paragraph(f"Figure {figure_number}. {cap}", style="Figure Caption Custom"); continue
        if s.startswith("\\noindent\\textbf{Keywords:}"):
            flush(); p=doc.add_paragraph(); r=p.add_run("Keywords: "); r.bold=True; p.add_run(clean(s.split("}",1)[1])); continue
        if s.startswith("\\textbf{") and ":}" in s:
            flush(); doc.add_paragraph(clean(s)); continue
        if not s:
            flush(); continue
        paragraph.append(s)
    flush()

    # Replace the LaTeX bibliography commands with the preformatted US Harvard list.
    doc.add_heading("References", level=1)
    refs = (PUB / "manuscript" / "references_harvard.md").read_text(encoding="utf-8").splitlines()[2:]
    for line in refs:
        if line.strip():
            p = doc.add_paragraph(line.strip()); p.paragraph_format.first_line_indent = Inches(-0.25); p.paragraph_format.left_indent = Inches(0.25); p.paragraph_format.space_after = Pt(5); p.style = doc.styles["Normal"]

    props = doc.core_properties
    props.title = "FairClaim-MM manuscript"; props.subject = "Fairness-aware calibrated insurance fraud triage"; props.author = "Author name(s) to be inserted"
    OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT); print(OUT)


if __name__ == "__main__":
    main()
